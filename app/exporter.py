#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""转写结果导出：TXT / SRT / Markdown / Word(docx)。"""
import os
import sys


def _docx_font() -> str:
    """Word 文档中文字体按平台选择（写死的字体在缺失时会回退默认）。"""
    return "PingFang SC" if sys.platform == "darwin" else "微软雅黑"


def _fmt_ts_srt(seconds: float) -> str:
    """SRT 时间戳：00:00:01,500"""
    ms = int(round(seconds * 1000))
    h, ms = divmod(ms, 3600000)
    m, ms = divmod(ms, 60000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def _fmt_ts_md(seconds: float) -> str:
    """Markdown 时间戳：00:01:23"""
    s = int(round(seconds))
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def _clean(segments: list) -> list:
    """归一化分段：去除文字首尾空白并丢弃空段（避免空行/空字幕）。"""
    out = []
    for seg in segments:
        text = (seg.get("text") or "").strip()
        if text:
            out.append({
                "start": seg.get("start", 0.0),
                "end": seg.get("end", 0.0),
                "text": text,
            })
    return out


def export_txt(segments: list, path: str, with_timestamps: bool = False) -> str:
    with open(path, "w", encoding="utf-8") as f:
        for seg in _clean(segments):
            if with_timestamps:
                f.write(f"[{_fmt_ts_md(seg['start'])} -> {_fmt_ts_md(seg['end'])}] ")
            f.write(seg["text"] + "\n")
    return path


def export_srt(segments: list, path: str) -> str:
    with open(path, "w", encoding="utf-8") as f:
        for i, seg in enumerate(_clean(segments), 1):
            f.write(f"{i}\n{_fmt_ts_srt(seg['start'])} --> {_fmt_ts_srt(seg['end'])}\n"
                    f"{seg['text']}\n\n")
    return path


def export_markdown(segments: list, path: str, title: str = "语音转写稿") -> str:
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# {title}\n\n")
        for seg in _clean(segments):
            f.write(f"- **[{_fmt_ts_md(seg['start'])}]** {seg['text']}\n")
    return path


def export_docx(segments: list, path: str, title: str = "语音转写稿") -> str:
    """生成简洁排版的 Word 文档（时间戳灰色小字 + 正文）。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    GRAY = RGBColor(0x7F, 0x7F, 0x7F)
    ACCENT = RGBColor(0x1F, 0x4E, 0x79)

    def set_cn_font(run, name=None):
        name = name or _docx_font()
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8); section.right_margin = Cm(2.8)

    p = doc.add_paragraph(); p.alignment = 1  # center
    r = p.add_run(title); r.bold = True
    r.font.size = Pt(18); r.font.color.rgb = ACCENT; set_cn_font(r)

    for seg in _clean(segments):
        p = doc.add_paragraph()
        ts = p.add_run(f"[{_fmt_ts_md(seg['start'])}] ")
        ts.font.size = Pt(8); ts.font.color.rgb = GRAY; ts.font.name = "Menlo"
        body = p.add_run(seg["text"]); body.font.size = Pt(11); set_cn_font(body)
        p.paragraph_format.space_after = Pt(4)

    doc.save(path)
    return path


def default_export_path(media_path: str, ext: str, out_dir: str = None) -> str:
    """默认导出路径：优先用户设置的导出目录，否则与媒体文件同目录、同名。"""
    base = os.path.splitext(os.path.basename(media_path))[0]
    directory = out_dir or os.path.dirname(media_path) or "."
    return os.path.join(directory, f"{base}_转写稿.{ext}")
